from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook


SUPPORTED_EXTENSIONS = {".txt", ".md", ".csv", ".docx", ".docm", ".pdf", ".xlsx", ".xlsm"}
TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "cp1250", "windows-1250", "iso-8859-2", "latin-1")


@dataclass(slots=True)
class LoadedDocument:
    path: Path
    text: str
    type: str
    metadata: dict[str, Any]


@dataclass(slots=True)
class TextChunk:
    text: str
    index: int
    paragraph_hint: str = ""


def file_sha256(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def classify_document(path: str | Path) -> dict[str, Any]:
    path = Path(path)
    name = path.name.lower()
    ext = path.suffix.lower()
    if "tvsz" in name:
        return {"type": "szabalyzat", "is_regulation": True, "regulation_name": "TVSZ"}
    if "tjsz" in name:
        return {"type": "szabalyzat", "is_regulation": True, "regulation_name": "TJSZ"}
    if "alapszab" in name or "szabalyzat" in name or "szabályzat" in name:
        return {"type": "szabalyzat", "is_regulation": True, "regulation_name": path.stem}
    if "jegyzokonyv" in name or "jegyzőkönyv" in name or "meeting" in name:
        return {"type": "jegyzokonyv", "is_regulation": False, "regulation_name": ""}
    if ext in {".xlsx", ".xlsm", ".csv"}:
        return {"type": "feladatlista", "is_regulation": False, "regulation_name": ""}
    if ext == ".pdf":
        return {"type": "pdf", "is_regulation": False, "regulation_name": ""}
    if ext in {".docx", ".docm"}:
        return {"type": "docx", "is_regulation": False, "regulation_name": ""}
    return {"type": "egyeb", "is_regulation": False, "regulation_name": ""}


def load_document(path: str | Path) -> LoadedDocument:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"A dokumentum nem talalhato: {path}")
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Nem tamogatott fajltipus: {ext}")

    if ext in {".txt", ".md"}:
        text, metadata = _read_text(path)
    elif ext == ".csv":
        text, metadata = _read_csv(path)
    elif ext in {".docx", ".docm"}:
        text, metadata = _read_docx(path)
    elif ext == ".pdf":
        text, metadata = _read_pdf(path)
    elif ext in {".xlsx", ".xlsm"}:
        text, metadata = _read_workbook(path)
    else:
        raise ValueError(f"Nem tamogatott fajltipus: {ext}")

    classification = classify_document(path)
    metadata.update(classification)
    metadata["extension"] = ext
    metadata["file_hash"] = file_sha256(path)
    return LoadedDocument(path=path, text=normalize_text(text), type=classification["type"], metadata=metadata)


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _decode_bytes(raw: bytes) -> tuple[str, str]:
    best_text = ""
    best_encoding = TEXT_ENCODINGS[-1]
    best_score = -1
    for encoding in TEXT_ENCODINGS:
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError:
            continue
        replacement_penalty = text.count("\ufffd") * 5
        hungarian_score = sum(text.count(char) for char in "áéíóöőúüűÁÉÍÓÖŐÚÜŰ")
        score = hungarian_score - replacement_penalty
        if score > best_score:
            best_text = text
            best_encoding = encoding
            best_score = score
    if not best_text:
        best_text = raw.decode("latin-1", errors="replace")
        best_encoding = "latin-1"
    return best_text, best_encoding


def _read_text(path: Path) -> tuple[str, dict[str, Any]]:
    text, encoding = _decode_bytes(path.read_bytes())
    return text, {"encoding": encoding}


def _read_csv(path: Path) -> tuple[str, dict[str, Any]]:
    text, encoding = _decode_bytes(path.read_bytes())
    sample = text[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample)
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    rendered = [" | ".join(cell.strip() for cell in row if cell is not None) for row in rows]
    return "\n".join(rendered), {"encoding": encoding, "rows": len(rows)}


def _read_docx(path: Path) -> tuple[str, dict[str, Any]]:
    document = Document(path)
    parts: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text.strip())
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                parts.append(" | ".join(values))
    return "\n".join(parts), {"paragraphs": len(document.paragraphs), "tables": len(document.tables)}


def _read_pdf(path: Path) -> tuple[str, dict[str, Any]]:
    try:
        import pdfplumber
    except ImportError as exc:
        raise RuntimeError("A PDF beolvasashoz telepiteni kell a pdfplumber csomagot.") from exc

    pages: list[str] = []
    with pdfplumber.open(path) as pdf:
        for index, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(f"[{index}. oldal]\n{text.strip()}")
    return "\n\n".join(pages), {"pages": len(pages)}


def _read_workbook(path: Path) -> tuple[str, dict[str, Any]]:
    workbook = load_workbook(path, read_only=False, data_only=True, keep_vba=path.suffix.lower() == ".xlsm")
    parts: list[str] = []
    for sheet in workbook.worksheets:
        parts.append(f"## Munkalap: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = [str(value).strip() for value in row if value not in (None, "")]
            if values:
                parts.append(" | ".join(values))
    return "\n".join(parts), {"sheets": [sheet.title for sheet in workbook.worksheets]}


def chunk_text(text: str, chunk_size: int = 1200, overlap: int = 180) -> list[TextChunk]:
    text = normalize_text(text)
    if not text:
        return []
    paragraphs = [item.strip() for item in text.split("\n\n") if item.strip()]
    chunks: list[TextChunk] = []
    current = ""
    current_hint = ""

    for paragraph in paragraphs:
        hint = extract_paragraph_hint(paragraph) or current_hint
        if len(current) + len(paragraph) + 2 <= chunk_size:
            current = f"{current}\n\n{paragraph}".strip()
            current_hint = hint
            continue
        if current:
            chunks.append(TextChunk(text=current, index=len(chunks), paragraph_hint=current_hint))
        current = paragraph
        current_hint = hint

        while len(current) > chunk_size:
            slice_text = current[:chunk_size]
            chunks.append(TextChunk(text=slice_text.strip(), index=len(chunks), paragraph_hint=current_hint))
            current = current[max(0, chunk_size - overlap) :]

    if current:
        chunks.append(TextChunk(text=current.strip(), index=len(chunks), paragraph_hint=current_hint))
    return chunks


def extract_paragraph_hint(text: str) -> str:
    match = re.search(r"(\d+\s*[./]?\s*§(?:\s*\([^)]+\))?)", text)
    if match:
        return re.sub(r"\s+", " ", match.group(1)).strip()
    return ""
