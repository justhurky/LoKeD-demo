from hok_ai.core.db import Database
from hok_ai.core.document_loader import load_document
from hok_ai.core.embeddings import DeterministicEmbeddingClient
from hok_ai.core.knowledge_base import KnowledgeBase
from hok_ai.core.meeting_parser import MeetingParser, normalize_meeting_json
from hok_ai.core.minutes_generator import MinutesGenerator
from hok_ai.training.dataset_builder import build_records, validate_output


def test_document_loader_reads_non_utf8_text(tmp_path):
    text = "Árvíztűrő tükörfúrógép EHÖK jegyzet"
    path = tmp_path / "notes_cp1250.txt"
    path.write_bytes(text.encode("cp1250"))

    loaded = load_document(path)

    assert text in loaded.text
    assert loaded.metadata["encoding"] in {"cp1250", "windows-1250"}


def test_document_loader_reads_docx(tmp_path):
    from docx import Document

    path = tmp_path / "meeting.docx"
    document = Document()
    document.add_paragraph("Kabinet meeting jegyzet.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Feladat"
    table.cell(0, 1).text = "Felelős"
    table.cell(1, 0).text = "TVSZ ellenőrzés"
    table.cell(1, 1).text = "Dani"
    document.save(path)

    loaded = load_document(path)

    assert "Kabinet meeting" in loaded.text
    assert "TVSZ ellenőrzés | Dani" in loaded.text


def test_chroma_knowledge_base_indexes_and_searches(tmp_path):
    db = Database(tmp_path / "hok_ai.sqlite")
    db.init()
    source = tmp_path / "TVSZ_szabalyzat.md"
    source.write_text("TVSZ 47. § A pótló vizsga tanulmányi ügyként kezelendő.", encoding="utf-8")
    kb = KnowledgeBase(db, embedding_client=DeterministicEmbeddingClient(), chroma_path=tmp_path / "chromadb")

    result = kb.index_file(source)
    hits = kb.search("TVSZ pótló vizsga", min_score=0.0)

    assert result["indexed"] == 1
    assert result["chunks"] >= 1
    assert hits
    assert hits[0].document_name == "TVSZ_szabalyzat.md"


def test_task_crud_and_dashboard(tmp_path):
    db = Database(tmp_path / "hok_ai.sqlite")
    db.init()

    task = db.create_task(
        {
            "title": "Kari elnökök értesítése",
            "owner": "Dani",
            "deadline": "2026-06-20",
            "priority": "sürgős",
            "source_sentence": "Dani küldje ki a kari elnököknek.",
        }
    )
    updated = db.update_task(task["id"], {"status": "Folyamatban", "owner": "Zalán"})

    assert updated["owner"] == "Zalán"
    assert db.dashboard()["counts"]["open_tasks"] == 1
    assert db.delete_task(task["id"]) is True
    assert db.task(task["id"]) is None


def test_people_memory_crud(tmp_path):
    db = Database(tmp_path / "hok_ai.sqlite")
    db.init()

    person = db.create_person({"name": "Aschi", "role": "logisztika", "aliases": "Ashi, Aschi"})
    updated = db.update_person(person["id"], {"committee": "Kabinet"})

    assert updated["committee"] == "Kabinet"
    assert "Ashi" in updated["aliases"]
    assert db.delete_person(person["id"]) is True


class FakeMeetingLLM:
    def chat(self, *_args, **_kwargs):
        return """
        {
          "feladatok": [{"cim": "Excel frissítés", "felelos": "Dani", "hataridő": "2026-06-20", "prioritas": "magas"}],
          "dontesek": [{"szoveg": "Elfogadtuk a bográcsjegy árát.", "datum": "2026-06-15", "szavazat": "igen"}],
          "kerdesek": [],
          "blokkatok": [{"szoveg": "Hiányzik a teremlista.", "megoldas_javaslat": "Kérjünk friss listát."}],
          "resztvevok": ["Dani"],
          "kaosz_index": 42
        }
        """


def test_meeting_parser_uses_structured_json():
    parser = MeetingParser(FakeMeetingLLM(), model="fake")

    result = parser.analyze("meeting jegyzet")

    assert result["feladatok"][0]["cim"] == "Excel frissítés"
    assert result["dontesek"]
    assert result["blokkatok"]
    assert result["kaosz_index"] == 42


def test_meeting_normalizer_accepts_internal_blocker_alias():
    result = normalize_meeting_json({"blokkolok": [{"szoveg": "Teszt"}]})

    assert result["blokkatok"][0]["szoveg"] == "Teszt"


class FakeMinutesLLM:
    def chat(self, messages, model, **kwargs):
        assert model == "hok-ai-jegyzokonyv-4b"
        assert "ELVART STRUKTURA" in messages[1]["content"]
        return "1. Rövid Összefoglaló:\nTeszt.\n\n2. Hivatalos Jegyzőkönyv:\nTeszt.\n\n3. Feladatnapló (Teendők):\nDani: Teszt - Határidő: Nem ismert\n\n4. Stratégiai Javaslatok:\nTeszt."


def test_minutes_generator_uses_hok_minutes_prompt():
    generator = MinutesGenerator(FakeMinutesLLM(), model="hok-ai-jegyzokonyv-4b")

    result = generator.generate("ööö hát akkor Dani intézi")

    assert "Feladatnapló" in result
    assert "Dani" in result


def test_training_dataset_builder_pairs_raw_and_gold(tmp_path):
    raw = tmp_path / "raw"
    gold = tmp_path / "gold"
    raw.mkdir()
    gold.mkdir()
    (raw / "001.txt").write_text("ööö Dani intézi a TVSZ ellenőrzést", encoding="utf-8")
    output = """1. Rövid Összefoglaló:
Teszt.

2. Hivatalos Jegyzőkönyv:
Teszt.

3. Feladatnapló (Teendők):
Dani: TVSZ ellenőrzés - Határidő: Nem ismert

4. Stratégiai Javaslatok:
Teszt.
"""
    (gold / "001.md").write_text(output, encoding="utf-8")

    records, issues = build_records(raw, gold)

    assert len(records) == 1
    assert not issues
    assert records[0]["input"].startswith("ööö")
    assert "System" not in records[0]["output"]


def test_training_dataset_validator_flags_missing_sections():
    issues = validate_output("1. Rövid Összefoglaló:\nTeszt")

    assert issues
